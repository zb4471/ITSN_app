from openpyxl import load_workbook
from openpyxl.drawing.image import Image

from PIL import Image as PILImage

class InvUpdater:
    def __init__(self):
        self.name = "InvUpdater"

    def update_inv(self, inv_excel, template_excel, stamp):

        wb = load_workbook(inv_excel)
        ws = wb.active

        # 1. 找到三个要复制的行
        targets = ["Gross Weight", "Packages", "Freight Charge"]
        target_rows = {}

        col_a_cells = [str(c.value).strip() if c.value else "" for c in ws["A"]]
        for tar in targets:
            row_idx = col_a_cells.index(tar) + 1
            target_rows[tar] = row_idx

        # 2. 定位标题行
        keys = ["INV NO", "物料编号", "报关品名", "数量", "单价"]

        header_cells = []
        for row in ws.iter_rows():
            values = [str(c.value).strip() if c.value else "" for c in row]
            if all(k in values for k in keys):
                header_cells = values
                break

        # 3. 确定“合同号”所在列
        if "合同号" in header_cells:
            col_inx = header_cells.index("合同号") + 1
            # 4. 如果“合同号”列还在，复制3个值
            for name, r in target_rows.items():
                src_cell = ws.cell(row=r, column=col_inx)
                dst_cell = ws.cell(row=r, column=col_inx + 1)
                dst_cell.value = src_cell.value 
            # 5. 删除“合同号”列
            ws.delete_cols(col_inx)

        # 6. 替换抬头
        wb_t = load_workbook(template_excel)
        ws_t = wb_t.active
        max_col = ws_t.max_column

        for r in range(2, 8):   # 2-7行
            for c in range(1, max_col + 1):
                v = ws_t.cell(row=r, column=c).value
                ws.cell(row=r, column=c).value = v
        
        wb_t.close()

        # 7. 添加贝克休斯章
        if not ws._images:  # sheet 没有图片才添加
            # ws.remove(ws._images[0])
            pil_img = PILImage.open(stamp)
            pixel_width, pixel_height = pil_img.size

            dpi = pil_img.info.get("dpi", (96, 96))
            dpi_x, dpi_y = dpi

            # 显示尺寸
            width_pt = pixel_width / dpi_x * 96
            height_pt = pixel_height / dpi_y * 96        
            
            # 位置列
            dst_row = target_rows["Freight Charge"] + 1

            img = Image(stamp)
            img.width = width_pt
            img.height = height_pt

            ws.add_image(img, f"D{dst_row}")


        wb.save(inv_excel)
        wb.close()       






if __name__ == "__main__":
    inv_updater = InvUpdater()

    inv_excel = r"D:\共享文件夹\金丽华\59-251017-004603 分拨数据 20251001-20251016\FBSN251059\ITSN2025110098发票信息.xlsx"
    template_excel = r"D:\共享文件夹\金丽华\抬头.xlsx"
    stamp = r"D:\共享文件夹\金丽华\贝克休斯章PNG.png"

    inv_updater.update_inv(inv_excel, template_excel, stamp)
