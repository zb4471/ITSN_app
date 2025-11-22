from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

def load_df_from_table(path, sheet_name, table_name=""):
    if not table_name:
        table_name = sheet_name

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    tbl = ws.tables[table_name]

    ref = tbl.ref
    range = ws[ref]
    data = [[cell.value for cell in row] for row in range]

    df = pd.DataFrame(data[1:], columns=data[0])
    return df


if __name__ == "__main__":
    path = Path(r"D:\共享文件夹\金丽华\参数表.xlsx")
    sheet_name = table_name = "接单字段映射表PCSZ"
    df = load_df_from_table(path, sheet_name, table_name)
    print(df)